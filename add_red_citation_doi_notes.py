from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor


PACKAGE = Path(
    "/Users/jxs794/Library/CloudStorage/OneDrive-UniversityofMiami/"
    "VLiSEMOD/Supplementary/Re-Submission-Package"
)
RED = RGBColor(0xC0, 0x00, 0x00)


def add_note(paragraph, note):
    if note in paragraph.text:
        return
    run = paragraph.add_run(" " + note)
    run.font.color.rgb = RED
    run.font.bold = True
    run.font.size = Pt(9)


def find_paragraph(doc, needle):
    for paragraph in doc.paragraphs:
        if needle in paragraph.text:
            return paragraph
    raise ValueError(f"Could not find paragraph: {needle}")


def note_clean_or_markup(path):
    doc = Document(path)
    add_note(
        find_paragraph(doc, "For a medicinal chemist"),
        "[ADD AFTER ‘...pocket.’ DOI: 10.1093/nar/28.1.235; 10.1038/s41573-021-00371-6; 10.1016/j.cbpa.2019.02.022; 10.1038/nchembio.2329. ADD AFTER ‘...biological activity.’ DOI: 10.1021/acs.jcim.0c00589; 10.1021/acs.jmedchem.8b01413; 10.1016/j.bmcl.2024.129676; 10.1038/s41467-018-08027-7.]",
    )
    add_note(
        find_paragraph(doc, "V-LiSEMOD complements the resources"),
        "[REPLACE STATIC NUMBERS WITH DOI: 10.1093/nar/28.1.235; 10.1093/bioinformatics/btu789; 10.1093/nar/gkad1004; 10.1093/nar/gkv315; 10.1016/j.jmb.2016.12.004; 10.12688/f1000research.7931.1; 10.1093/nar/gkae768; 10.1093/nar/gkae1091. CLARIFICATION DOI: 10.1093/nar/gkv315; 10.1021/ci200227u; 10.1016/j.jmb.2016.12.004.]",
    )
    add_note(
        find_paragraph(doc, "Static co-crystal structures provide"),
        "[REMOVE SOUTHAN ONLY: DOI 10.1016/j.coph.2016.07.002.]",
    )
    add_note(
        find_paragraph(doc, "Two published HIV-1 protease series"),
        "[REPLACE STATIC 45,46 WITH DOI: 10.1021/acs.bioconjchem.9b00195; 10.1111/j.1747-0285.2009.00921.x.]",
    )
    add_note(
        find_paragraph(doc, "For each retained ligand"),
        "[USE RDKit DOI: 10.5281/zenodo.19922430. REMOVE OBSOLETE DOI: 10.5281/zenodo.591637.]",
    )
    doc.save(path)


def note_s5(path):
    doc = Document(path)
    table = doc.tables[5]
    red_notes = {
        1: " [FINAL CITE DOI: 10.1093/nar/28.1.235; 10.1093/bioinformatics/btu789; 10.1093/nar/gkae1091.]",
        2: " [FINAL CITE DOI: 10.1093/nar/gkv315; 10.1016/j.jmb.2016.12.004.]",
        3: " [FINAL CITE DOI: 10.12688/f1000research.7931.1.]",
        4: " [FINAL CITE DOI: 10.1093/nar/gkad1004.]",
        5: " [FINAL CITE DOI: 10.1093/nar/gkae768.]",
    }
    for row_index, note in red_notes.items():
        add_note(table.rows[row_index].cells[0].paragraphs[0], note)
    examples = doc.tables[6]
    for row_index, note in {
        1: " [FINAL CITE DOI: 10.1021/acs.bioconjchem.9b00195.]",
        2: " [FINAL CITE DOI: 10.1111/j.1747-0285.2009.00921.x.]",
        3: " [FINAL CITE DOI: 10.1038/s41589-020-00689-z; REMOVE TEMPORARY DOI MARKER.]",
    }.items():
        add_note(examples.rows[row_index].cells[2].paragraphs[0], note)
    doc.save(path)


def note_response(path):
    doc = Document(path)
    add_note(
        find_paragraph(doc, "We removed the residual citation group"),
        "[REVISE: Southan DOI 10.1016/j.coph.2016.07.002 removed; directly relevant structural/PROTAC literature restored.]",
    )
    add_note(
        find_paragraph(doc, "generic RDKit concept DOI has been flagged"),
        "[FINAL RDKit DOI: 10.5281/zenodo.19922430.]",
    )
    doc.save(path)


def main():
    note_clean_or_markup(PACKAGE / "ACS_Infectious_Disease_2026_Schulz_etAL_Revised_MENDELEY_REPAIRED.docx")
    note_clean_or_markup(PACKAGE / "ACS_Infectious_Disease_2026_Schulz_etAL_MARKUP_MENDELEY_REPAIRED.docx")
    note_response(PACKAGE / "Response_to_Reviewer_Comments_MENDELEY_REPAIRED.docx")
    note_s5(PACKAGE / "Supplementary_File_5_Supplemental_MENDELEY_REPAIRED.docx")


if __name__ == "__main__":
    main()
