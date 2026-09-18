from __future__ import annotations

import csv
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


SRC = Path('/Users/jxs794/Library/CloudStorage/OneDrive-UniversityofMiami/VLiSEMOD/Supplementary')
OUT = SRC / 'Re-Submission-Package'
ORIGINAL = Path('/Users/jxs794/Library/CloudStorage/OneDrive-UniversityofMiami/VLiSEMOD/ACS_Infectious_Disease_2026_Schulz_etAL.docx')
BLUE = RGBColor(0x00, 0x00, 0xFF)


def normalize(text: str) -> str:
    return ' '.join((text or '').split())


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        if child.tag.endswith('}pPr'):
            continue
        p.remove(child)


def set_paragraph_text(paragraph, text, color=None, bold=False, size=None):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return run


def find_paragraph(doc, needle):
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise ValueError(f'paragraph not found: {needle}')


def add_before(anchor, text, style=None, bold=False, color=None, size=None):
    new_p = OxmlElement('w:p')
    anchor._p.addprevious(new_p)
    p = Paragraph(new_p, anchor._parent)
    if style:
        p.style = style
    set_paragraph_text(p, text, color=color, bold=bold, size=size)
    return p


def set_cell_text(cell, text, size=6, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_margins(cell, top=45, start=45, bottom=45, end=45):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(f'{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{m}')
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w', str(v))
        node.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type', 'dxa')


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'true')
    trPr.append(tblHeader)


def add_table(doc, headers, rows, widths, font_size=6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.autofit = False
    set_repeat_header(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        set_cell_text(cell, h, font_size, bold=True)
        set_cell_margins(cell)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            set_cell_text(cells[i], value, font_size)
            set_cell_margins(cells[i])
    return table


def make_clean():
    source = SRC / 'ACS_Infectious_Disease_2026_Schulz_etAL_Reviewer_Pass12_RESUBMISSION_FINAL.docx'
    target = OUT / 'ACS_Infectious_Disease_2026_Schulz_etAL_Revised_R2_FINAL.docx'
    shutil.copy2(source, target)
    doc = Document(target)

    p = find_paragraph(doc, 'For a medicinal chemist, the practical question')
    set_paragraph_text(p, 'For a medicinal chemist, the practical question posed by a viral co-crystal structure is not simply what binds, but what can be changed without sacrificing the recognition features that support binding. This question becomes especially important when a known inhibitor is being converted into an analogue, conjugate or degrader warhead. A useful design workflow must identify the ligand features that should be preserved, the atoms or substituents that remain accessible to solvent and the spatial route by which new chemistry could extend away from the pocket. Accordingly, V-LiSEMOD looks for positions that make few important protein contacts, remain exposed to solvent, are chemically tractable and point away from the binding pocket. These criteria reflect standard medicinal-chemistry reasoning and are used to generate testable hypotheses; they were not fitted to a training set of known tolerated sites. Agreement across related co-crystal structures increases confidence in the hypothesis but does not establish that modification will preserve affinity or biological activity. [REMOVE CURRENT CITATION HERE]')

    p = find_paragraph(doc, 'V-LiSEMOD complements the resources')
    set_paragraph_text(p, p.text + ' PLIP and LigPlot+ are cited here as complementary interaction-analysis and visualization resources; the protein–ligand contact records reported and analyzed in V-LiSEMOD were generated with pdbe-arpeggio 1.4.4.')

    p = find_paragraph(doc, 'These comparisons nominate regions to test')
    set_paragraph_text(p, p.text + ' Across the 12 occurrence-resolved Figure 5 DR7 structures, all 51 common mapped positions had mapped evidence in all 12 occurrences. Supplemental Table S5.8 reports position-level protein-contact and solvent-exposure frequencies, mean protein-contact records and standard deviations; for example, positions 29/CAO and 30/CAS were exposed in all 12 structures and had the lowest mean protein-contact burden (5.00 records each), while the table also identifies positions with greater cross-structure variability. These are cross-structure observations, not independent replicates.')

    p = find_paragraph(doc, 'For each retained ligand, V-LiSEMOD resolved chemical identity')
    set_paragraph_text(p, p.text + ' RDKit 2026_03_2 was used for the molecular and SMARTS operations described here. [DOI TO ADD MANUALLY: 10.5281/zenodo.19922430]')

    p = find_paragraph(doc, 'Mapping QC covered all 7,335 ligand instances')
    set_paragraph_text(p, 'Mapping QC covered all 7,335 ligand instances. Complete maps were obtained for 3,864 instances, with another 570 becoming complete after deterministic alternate-conformer selection. A further 2,182 were valid partial maps because the deposited atom set did not fully match the chemical reference; only supported atom correspondences were retained for these cases. The remaining 719 records are explicitly retained as QC-controlled pending-remediation cases and are not downstream mapping eligible. We did not manually reclassify ambiguous, partial or pending-remediation records as complete, and analyses that require a complete common-atom map exclude instances without sufficient correspondence. Supplemental Table S5.10 gives the release-level category counts and percentages.')

    doc.save(target)
    return target


def make_s5():
    source = SRC / 'Supplementary_File_5_Supplemental_RESUBMISSION_FINAL.docx'
    target = OUT / 'Supplementary_File_5_Supplemental_R2_FINAL.docx'
    shutil.copy2(source, target)
    doc = Document(target)
    p = find_paragraph(doc, 'This file provides additional database')
    set_paragraph_text(p, p.text.replace('ACS Infectious Diseases submission', 'ACS Omega submission').replace('Supplemental Tables S5.6 and S5.7 compare V-LiSEMOD with complementary resources and summarize published HIV-1 protease modification examples, respectively.', 'Supplemental Tables S5.6–S5.10 compare V-LiSEMOD with complementary resources, summarize qualitative retrospective examples, and report Reviewer 2 quantitative cross-structure, heuristic-definition and mapping-QC audits.'))

    # Add the peer-reviewed SARS-CoV-2 example to existing S5.7.
    t = doc.tables[6]
    cells = t.add_row().cells
    values = [
        'Structure-resolved conjugation example',
        'PDB 6Z2E; ligand Q5T; SARS-CoV-2 Mpro. This covalent activity-based probe is present in the V-LiSEMOD ligand-instance inventory.',
        'Rut et al. reported Biotin-PEG(4)-Abu-Tle-Leu-Gln-vinylsulfone (B-QS1-VS), determined its 1.7 Å Mpro co-crystal structure, and used related biotin- and fluorophore-bearing probes for Mpro detection. The parent inhibitor and probes retained measurable target engagement; the paper emphasizes that the flexible label tail contacts a neighboring dimer in this crystal lattice, so the structure is interpreted cautiously as a qualitative extension example rather than a solution-state linker model. [DOI TO ADD MANUALLY: 10.1038/s41589-020-00689-z]',
        'Provides a second, directly viral and structure-resolved example in which a biotin-PEG(4) extension was accommodated while the viral-protease recognition region remained engaged.',
        'Qualitative only. Current V-LiSEMOD mapping for Q5T is not sufficient for automatic atom-level attachment ranking, so this row does not claim that the release automatically flags the same extension site.'
    ]
    for cell, value in zip(cells, values):
        set_cell_text(cell, value, size=7)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = Inches(0.45)
    sec.bottom_margin = Inches(0.45)
    sec.left_margin = Inches(0.45)
    sec.right_margin = Inches(0.45)

    p = doc.add_paragraph('Supplemental Table S5.8. Cross-structure consistency of mapped DR7 positions across the Figure 5 HIV-1 protease structures.')
    p.runs[0].bold = True
    doc.add_paragraph('The selected occurrences are 2AQU B/300, 2FXD A/102, 2FXE A/102, 2O4K A/301, 3EKW B/100, 3EKY A/100, 3EL1 A/100, 3EL9 A/100, 3EM4 A/100, 3EM4 V/100, 3OXX A/100 and 3OXX C/100. The denominator is the number of structures with valid mapped evidence for the position; missing mapping/evidence is not treated as zero. Protein-contact records are occurrence-resolved rows in the release protein-contact table. SD is the sample standard deviation across structures with mapped evidence.')
    with open('/tmp/dr7_protein_summary.csv', newline='') as f:
        table_rows = list(csv.reader(f))[1:]
    add_table(doc, ['SMILES\nindex', '3EKY DR7\natom', 'Mapped\nevidence', 'Protein-contact\nfrequency', 'Solvent-exposure\nfrequency', 'Mean protein-contact\nrecords', 'SD'], table_rows, [0.65, 0.85, 1.0, 1.55, 1.55, 1.55, 0.55], 5.8)

    p = doc.add_paragraph('Supplemental Table S5.9. V-LiSEMOD PROTACability and attachment-site heuristic definitions.')
    p.runs[0].bold = True
    doc.add_paragraph('All entries below are transparent heuristic prioritization rules. They were not trained, fitted, or experimentally calibrated against a degrader-success dataset and do not predict degradation accuracy, optimal linker length, E3 compatibility, ternary-complex success, or ubiquitination. “Current” refers to the version-pinned v2.8 release implementation; legacy compatibility columns are not treated as active score components.')
    s59 = [
        ['Candidate attachment atom', 'Non-H atom; mapped; SASA >0.1 Å²; strong-contact count.', 'candidate = mapped AND solvent-exposed AND strong-contact count = 0.', 'binary', 'No tier.', 'Screened ligand-side evidence set; not a synthetic feasibility call.'],
        ['Attachment priority', 'Mapping, exposure, strong contacts, unique protein partners, chemical context, outward score, local corridor.', '20 mapped + 30 exposed + 20 no strong contacts + partner term (10/7/3/0) + chemistry term (10/6/3/1) + outward term (8 or 5) + 7 for outward and locally clear; non-candidates capped at 49.', '0–100', 'High ≥80; moderate ≥60; exploratory ≥40; low <40.', 'High priority additionally requires candidate core, direct/conditional chemical support, outward orientation and local clearance.'],
        ['Interaction-preservation subscore', 'Mapped, exposed atoms; candidate-core count; strong-contact atoms.', '100 × candidate-core atoms / mapped exposed atoms; subtract 30 if all are strong-contact atoms and none candidate, otherwise subtract min(20, 4 × strong-contact atoms).', '0–100', 'No displayed tier.', 'Warhead-linkability component only; descriptive structural signal.'],
        ['Warhead linkability', 'Ligand context, SMILES/RDKit validity, mapped atoms, exposed atoms, functional-group types, candidate-core count, interaction-preservation subscore.', '10 context + 12 SMILES + 8 valid RDKit + min(15,5+0.5×mapped) + exposure term [min(20,8+3×mapped exposed), or partial-map alternative] + min(15,5+2×FG types) + candidate term (20/16/10/0) + 0.10×preservation; explicit missing-data penalties; clamped.', '0–100', 'High ≥75; moderate ≥55; exploratory ≥35; weak <35.', 'Ligand-centered evidence only; target lysines evaluated separately.'],
        ['Ligand Exit-Vector Cue', 'Screened candidate attachment atoms; ligand/pocket centroids; protein atoms in forward corridor.', '100 × [0.40 × (outward candidates/candidates) + 0.60 × (outward AND locally clear candidates/candidates)]. Outward is ligand-centroid→atom relative to ligand-centroid→pocket-centroid; local corridor is 8 Å forward, 2 Å radius, with no protein atom.', '0–100', 'Label: locally clear if any outward/clear candidate; otherwise outward-obstructed, no-outward, or no-candidate.', 'Local ligand-extension cue. Clear means clear to the sampled 8 Å cap, not unlimited clearance.'],
        ['Target lysine accessibility / structural priority', 'All chain lysines; observed NZ atoms; NZ SASA >1 and >5 Å².', '10 × (observed NZ/all Lys) + 60 × (NZ SASA >1/all Lys) + 30 × (NZ SASA >5/all Lys); missing NZ lowers score rather than being called buried.', '0–100', 'Structural priority: high ≥70; moderate ≥45; low <45.', 'Target-surface cue; ligand-to-lysine distance is not scored.'],
        ['Degrader-design readiness', 'Warhead-linkability score and target-lysine-accessibility score.', '0.60 × warhead linkability + 0.40 × target lysine accessibility; if no NZ SASA >1 Å², cap at 34.99. Ligand Exit-Vector Cue is displayed and contributes upstream to ligand evidence but is not a separate readiness weight.', '0–100', 'High ≥75; moderate ≥55; exploratory ≥35; weak <35.', 'Triage for later explicit modeling and experiments only. Legacy ternary_geometry_cue_score is retained as 0 for compatibility and is not used by current v2.8 readiness.'],
    ]
    add_table(doc, ['Score / cue', 'Inputs', 'Rule / equation', 'Range', 'Tier / threshold', 'Interpretation boundary'], s59, [1.25, 1.55, 3.15, 0.65, 1.35, 1.55], 5.4)

    p = doc.add_paragraph('Supplemental Table S5.10. Release-level ligand atom-mapping QC.')
    p.runs[0].bold = True
    s510 = [
        ['Complete directly', '3,864', '52.68%', 'Complete supported mapping without alternate-conformer resolution.'],
        ['Complete after deterministic alternate-conformer selection', '570', '7.77%', 'Complete supported mapping after coherent conformer selection; not manual correction.'],
        ['Valid partial map', '2,182', '29.75%', 'Only supported correspondences retained; not complete.'],
        ['QC-controlled pending remediation', '719', '9.80%', 'No downstream mapping eligibility in the frozen release; includes records deferred for explicit remediation rather than coerced to complete.'],
        ['Total resolved-chemistry mapping population', '7,335', '100.00%', 'Release denominator.'],
    ]
    add_table(doc, ['Mapping outcome', 'Count', 'Percent of 7,335', 'Interpretation'], s510, [2.2, 0.8, 1.2, 5.2], 6.5)
    doc.save(target)
    return target


def make_response():
    source = SRC / 'VLiSEMOD_Response_to_Reviewer_Comments_RESUBMISSION_FINAL.docx'
    target = OUT / 'Response_to_Reviewer_Comments_R2_FINAL.docx'
    shutil.copy2(source, target)
    doc = Document(target)
    anchor = find_paragraph(doc, 'Closing statement')
    style = anchor.style
    add_before(anchor, 'Reviewer 2', style=style, bold=True)
    entries = [
        ('Comment 1. Cross-structure quantitative consistency.', 'Please provide a quantitative summary across the 12 atazanavir/DR7 HIV-1 protease structures in Figure 5, including mapped-position contact frequency, solvent-exposure frequency, mean contact count and variability, with missing mappings distinguished from zero.', 'Response. We added Supplemental Table S5.8, an occurrence-resolved 51-position DR7 table across the 12 Figure 5 structures (2AQU B/300; 2FXD A/102; 2FXE A/102; 2O4K A/301; 3EKW B/100; 3EKY A/100; 3EL1 A/100; 3EL9 A/100; 3EM4 A/100 and V/100; 3OXX A/100 and C/100). For each common SMILES position it reports structures with valid mapped evidence, structures/fraction with one or more protein-contact records, structures/fraction solvent exposed, mean protein-contact records and sample SD. Missing mapping/evidence is excluded from the position denominator rather than counted as zero. All 51 DR7 positions were mapped in all 12 selected occurrences. The new Results summary identifies persistently exposed, relatively lower-contact positions (for example, 29/CAO and 30/CAS) without treating related structures as independent replicates. This is a descriptive consistency audit, not evidence that any position tolerates modification.'),
        ('Comment 2. PROTACability score definitions.', 'Please provide equations or decision rules, weights, normalization, thresholds, tier cutoffs, missing-data behavior and an explicit statement of whether the displayed scores are calibrated.', 'Response. We added Supplemental Table S5.9, which traces the current version-pinned v2.8 implementation for candidate attachment atoms, attachment priority, interaction-preservation, warhead linkability, the Ligand Exit-Vector Cue, target lysine accessibility/structural priority and degrader-design readiness. The table gives inputs, exact rules or equations, score ranges, tier cutoffs and missing-data behavior. It preserves the current cue definition: screened candidates only; ligand-centroid-to-atom direction relative to ligand-centroid-to-pocket-centroid direction; an 8 Å forward, 2 Å-radius protein-free corridor for local clearance; and a 40% outward plus 60% outward-and-locally-clear score. We explicitly state that these are transparent heuristic prioritization scores, not trained, fitted or experimentally calibrated predictors of degradation, linker length, E3 compatibility, ternary-complex success or ubiquitination. Legacy compatibility fields are identified rather than documented as active current score components.'),
        ('Comment 3. Mapping QC.', 'Please provide release-level quantitative QC for ligand mapping, define the categories and state how ambiguous or partial mappings were handled.', 'Response. We retained the manuscript mapping-QC paragraph and added Supplemental Table S5.10 for auditability. The frozen release verifies 3,864/7,335 complete maps (52.68%), 570/7,335 complete after deterministic alternate-conformer selection (7.77%), 2,182/7,335 valid partial maps (29.75%), and 719/7,335 QC-controlled pending-remediation cases (9.80%). The final release represents these remaining cases as pending remediation rather than the earlier preliminary adapter/timeout/failure subdivision. No ambiguous, partial or pending-remediation map was manually forced into the complete category; analyses needing complete common-atom correspondence exclude insufficiently mapped instances.'),
        ('Comment 4. Retrospective medicinal-chemistry examples.', 'Please provide two published tolerated modification/conjugation examples and one negative poor-tolerance case.', 'Response. Supplemental Table S5.7 now contains two positive qualitative examples and one negative comparator. The existing HIV-1 protease B0F/6IXD example reports direct biotinylation at a solvent-facing appendage with retained nanomolar inhibitory activity. We added the structure-resolved SARS-CoV-2 Mpro Q5T/6Z2E biotin-PEG(4) activity-based probe as a second viral-ligand extension example, explicitly noting its crystal-lattice and current atom-mapping limitations. The KNI series in 3KDB/3KDC/3KDD remains the negative buried-pocket steric-limit example. These rows are literature-grounded qualitative checks only and are not a calibration or three-case predictive benchmark.'),
        ('Comment 5a. Southan reference.', 'The Southan GPCR-database reference appears weakly related to the opening modification-site-design argument.', 'Response. We removed the residual citation group at the affected opening-design sentence from the working manuscript and inserted the required editorial handoff marker, [REMOVE CURRENT CITATION HERE], rather than preserving Southan as support for the statement. The reference list was not edited or renumbered. This marker directs the authors to remove the inappropriate citation-manager entry while retaining or restoring only any directly relevant support they judge necessary.'),
        ('Comment 5b. PLIP/LigPlot+ clarification.', 'Clarify that PLIP or LigPlot+ did not generate the reported V-LiSEMOD contact calculations.', 'Response. We added one concise Introduction clarification: PLIP and LigPlot+ are cited as complementary interaction-analysis and visualization resources, whereas the protein–ligand contact records reported and analyzed in V-LiSEMOD were generated with pdbe-arpeggio 1.4.4.'),
        ('Comment 5c. RDKit version-specific DOI.', 'The generic RDKit concept DOI does not uniquely identify RDKit 2026_03_2.', 'Response. The generic RDKit concept DOI has been flagged for replacement with the version-specific DOI corresponding to RDKit 2026_03_2: 10.5281/zenodo.19922430. The exact temporary marker was added to the RDKit Methods sentence; the reference list was not edited.'),
        ('Comment 5d. Journal name in S5.', 'Correct inherited wording that describes this as an ACS Infectious Diseases submission.', 'Response. We corrected the S5 introductory submission wording to ACS Omega and searched the final candidate package for incorrect ACS Infectious Diseases statements. Legacy physical filenames were preserved because they are filenames, not scientific or submission statements.'),
    ]
    for heading, comment, response in entries:
        add_before(anchor, heading, bold=True)
        add_before(anchor, comment)
        add_before(anchor, response)
        if heading.startswith('Comment 1'):
            loc = 'Revised manuscript/SI location: Results Section 2.5; Supplemental Table S5.8.'
        elif heading.startswith('Comment 2'):
            loc = 'Revised manuscript/SI location: Results Section 2.6; Supplemental Table S5.9.'
        elif heading.startswith('Comment 3'):
            loc = 'Revised manuscript/SI location: Methods Section 5.6; Supplemental Table S5.10.'
        elif heading.startswith('Comment 4'):
            loc = 'Revised manuscript/SI location: Supplemental Table S5.7.'
        elif heading.startswith('Comment 5a'):
            loc = 'Revised manuscript/SI location: Introduction opening modification-site paragraph.'
        elif heading.startswith('Comment 5b'):
            loc = 'Revised manuscript/SI location: Introduction resource-comparison paragraph.'
        elif heading.startswith('Comment 5c'):
            loc = 'Revised manuscript/SI location: Methods Section 5.5, RDKit sentence.'
        else:
            loc = 'Revised manuscript/SI location: Supplementary File S5 introductory paragraph.'
        add_before(anchor, loc)
    doc.save(target)
    return target


def make_cover():
    source = SRC / 'Re-Submission-Package/CoverLETTER_RESUBMISSION_FINAL.docx'
    target = OUT / 'CoverLETTER_RESUBMISSION_R2_FINAL.docx'
    shutil.copy2(source, target)
    doc = Document(target)
    p = find_paragraph(doc, 'The revised Supporting Information comprises')
    set_paragraph_text(p, 'The revised Supporting Information comprises Supplementary File 1, Database Schema (Supplementary_File_1_Database_Schema.docx); Supplementary File 2, Programmatic Manifest (Supplementary_File_2_API_Manifest.docx); Supplementary File 3, Programmatic Access Guide (Supplementary_File_3_Programmatic_Access_Guide.docx); Supplementary File 4, machine-readable manifest v1.0 (physical filename: Supplementary_File_4_API_Manifest_v0_1.txt); and Supplementary File 5, Supplemental Material (Supplementary_File_5_Supplemental.docx). The matching Supplementary_File_4_API_Manifest_v0_1.json is retained with the project/repository release; the corresponding plain-text manifest is supplied for the journal submission.')
    doc.save(target)
    return target


def mark_changed(final_path):
    target = OUT / 'ACS_Infectious_Disease_2026_Schulz_etAL_MARKUP_R2_FINAL.docx'
    shutil.copy2(final_path, target)
    final = Document(target)
    original = Document(ORIGINAL)
    original_paragraphs = {normalize(p.text) for p in original.paragraphs if normalize(p.text)}
    original_tables = {normalize(' | '.join(c.text for r in t.rows for c in r.cells)) for t in original.tables}
    title = final.paragraphs[0]
    marker = add_before(title, 'Supporting Information for Review Only — Marked Revised Manuscript (blue text marks revised or added material relative to the originally reviewed manuscript).', bold=True, color=BLUE, size=9)
    for p in final.paragraphs:
        text = normalize(p.text)
        if text and text not in original_paragraphs and not text.startswith('References'):
            for r in p.runs:
                r.font.color.rgb = BLUE
    for table in final.tables:
        signature = normalize(' | '.join(c.text for r in table.rows for c in r.cells))
        if signature not in original_tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.color.rgb = BLUE
    final.save(target)
    return target


if __name__ == '__main__':
    OUT.mkdir(exist_ok=True)
    clean = make_clean()
    s5 = make_s5()
    response = make_response()
    cover = make_cover()
    markup = mark_changed(clean)
    for path in (clean, markup, response, s5, cover):
        print(path)
