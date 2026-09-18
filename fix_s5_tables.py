from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(
    "/Users/jxs794/Library/CloudStorage/OneDrive-UniversityofMiami/"
    "VLiSEMOD/Supplementary/Re-Submission-Package/"
    "Supplementary_File_5_Supplemental.docx"
)
OUTPUT = SOURCE.with_name("Supplementary_File_5_Supplemental_FORMATTED.docx")


def set_cell_margins(cell, top=55, start=65, bottom=55, end=65):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is None:
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 0.95
                for run in paragraph.runs:
                    run.font.size = Pt(7.2)

    # Repeat the column labels if Word/LibreOffice moves any body rows forward.
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    if header_pr.find(qn("w:tblHeader")) is None:
        header_pr.append(OxmlElement("w:tblHeader"))


def main():
    doc = Document(SOURCE)

    # The final section was tagged landscape but retained portrait dimensions.
    # Restore true Letter landscape so the existing tables fit inside its margins.
    final_section = doc.sections[-1]
    final_section.orientation = WD_ORIENTATION.LANDSCAPE
    final_section.page_width = Inches(11)
    final_section.page_height = Inches(8.5)
    final_section.left_margin = Inches(0.55)
    final_section.right_margin = Inches(0.55)
    final_section.top_margin = Inches(0.55)
    final_section.bottom_margin = Inches(0.55)

    # Maintain the existing information, but use widths that total less than the
    # usable 9.9-inch text block and reserve room for readable wrapped prose.
    set_table_width(doc.tables[8], [1.18, 1.45, 3.05, 0.60, 1.32, 1.50])
    set_table_width(doc.tables[9], [2.05, 0.78, 1.26, 5.00])

    # Keep headings legible after the table adjustment.
    for idx in (18, 20):
        p = doc.paragraphs[idx]
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(4)

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
