"""Build and insert the V-LiSEMOD v1.0 logical schema connection map."""
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor


DOCX = Path(
    "/Users/jxs794/Library/CloudStorage/OneDrive-UniversityofMiami/"
    "VLiSEMOD/Supplementary/Supplementary_File_1_Database_Schema.docx"
)
PNG = Path("/tmp/vlisemod_s1_v1_schema_map.png")
PNG_B = Path("/tmp/vlisemod_s1_v1_schema_map_expanded.png")
PNG_B1 = Path("/tmp/vlisemod_s1_v1_schema_map_detail_core.png")
PNG_B2 = Path("/tmp/vlisemod_s1_v1_schema_map_detail_evidence.png")
PNG_B3 = Path("/tmp/vlisemod_s1_v1_schema_map_detail_protac.png")

W, H = 2800, 2200
BG = "#FFFFFF"
NAVY = "#163A63"
TEXT = "#15202B"
MUTED = "#536270"
GRID = "#9EABB7"


def font(size, bold=False):
    paths = (
        ["/Library/Fonts/Arial Bold.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"]
        if bold
        else ["/Library/Fonts/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf"]
    )
    for path in paths:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


F_TITLE = font(57, True)
F_SUB = font(30)
F_BAND = font(30, True)
F_NODE = font(27, True)
F_BODY = font(24)
F_SMALL = font(22)
F_MINI = font(30)
F_MINI_BOLD = font(32, True)


def wrap(draw, text, fnt, width):
    words, lines, line = text.split(), [], ""
    for word in words:
        candidate = f"{line} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=fnt)[2] <= width or not line:
            line = candidate
        else:
            lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def box(draw, xy, title, body, fill, accent):
    x0, y0, x1, y1 = xy
    draw.rounded_rectangle(xy, radius=26, fill=fill, outline=accent, width=5)
    draw.rounded_rectangle((x0, y0, x1, y0 + 48), radius=22, fill=accent)
    draw.rectangle((x0, y0 + 25, x1, y0 + 48), fill=accent)
    title_lines = wrap(draw, title, F_NODE, x1 - x0 - 48)
    y = y0 + 70
    for line in title_lines:
        draw.text((x0 + 24, y), line, font=F_NODE, fill=TEXT)
        y += 34
    y += 10
    for line in wrap(draw, body, F_BODY, x1 - x0 - 48):
        draw.text((x0 + 24, y), line, font=F_BODY, fill=MUTED)
        y += 31


def arrow(draw, start, end, color=NAVY, width=8):
    x0, y0, x1, y1 = *start, *end
    draw.line((x0, y0, x1, y1), fill=color, width=width)
    dx, dy = x1 - x0, y1 - y0
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    tip = (x1, y1)
    p1 = (x1 - ux * 25 + px * 13, y1 - uy * 25 + py * 13)
    p2 = (x1 - ux * 25 - px * 13, y1 - uy * 25 - py * 13)
    draw.polygon((tip, p1, p2), fill=color)


def build_png():
    im = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(im)
    draw.text((90, 65), "V-LiSEMOD v1.0 logical schema connection map", font=F_TITLE, fill=NAVY)
    draw.text(
        (92, 142),
        "Physical application tables grouped by role; arrows denote validated logical/provenance joins, not SQLite foreign keys.",
        font=F_SUB,
        fill=MUTED,
    )

    groups = [
        (90, 225, 695, 250, "1  STRUCTURE CONTEXT", "#E8F1FB", "#3F7FBF"),
        (755, 225, 1360, 250, "2  LIGAND IDENTITY & INSTANCES", "#E9F7F5", "#21897E"),
        (1420, 225, 2025, 250, "3  EVIDENCE GENERATION", "#FFF4E5", "#C77B1A"),
        (2085, 225, 2710, 250, "4  TARGET & ATTACHMENT", "#F4ECFA", "#8751A1"),
    ]
    for x0, y0, x1, _, label, fill, accent in groups:
        draw.rounded_rectangle((x0, y0, x1, y0 + 56), radius=18, fill=accent)
        draw.text((x0 + 20, y0 + 11), label, font=F_BAND, fill="white")

    # Primary data spine
    box(draw, (90, 320, 695, 590), "structures", "Unique deposited PDB/mmCIF entry", "#E8F1FB", "#3F7FBF")
    box(draw, (90, 650, 695, 970), "structure_classifications + structure_context", "Frozen source path and normalized virus/target context", "#E8F1FB", "#3F7FBF")
    box(draw, (755, 320, 1360, 590), "ligands + ligand_chemistry_sources", "Chemical-component identity and versioned chemistry candidates", "#E9F7F5", "#21897E")
    box(draw, (755, 650, 1360, 970), "ligand_instances", "Deposited ligand occurrence: structure, model, chain, component, residue, insertion code", "#E9F7F5", "#21897E")
    box(draw, (755, 1030, 1360, 1295), "ligand_instance_atoms", "Occurrence-resolved ligand atoms; the atom-level join key", "#E9F7F5", "#21897E")

    box(draw, (1420, 320, 2025, 620), "mapping + SASA", "ligand_mapping_runs; ligand_smiles_atom_mapping; ligand_sasa_atoms", "#FFF4E5", "#C77B1A")
    box(draw, (1420, 680, 2025, 980), "contacts + functional groups", "ligand_arpeggio_runs; arpeggio_*; ligand_functional_group_*; ligand_binding_pocket_atoms", "#FFF4E5", "#C77B1A")
    box(draw, (1420, 1040, 2025, 1295), "geometry", "ligand_atom_geometry and evidence-generation run provenance", "#FFF4E5", "#C77B1A")

    box(draw, (2085, 320, 2710, 615), "canonical targets + lysines", "canonical_ligand_targets; target_surface_lysines; target_chain_geometry", "#F4ECFA", "#8751A1")
    box(draw, (2085, 675, 2710, 975), "PROTACability context", "protacability_target_context; ligand_inventory; lysine_proximity; warhead_linkability; degrader_readiness", "#F4ECFA", "#8751A1")
    box(draw, (2085, 1035, 2710, 1295), "attachment-site assessment", "protacability_attachment_sites and protacability_attachment_site_summary", "#F4ECFA", "#8751A1")

    # Horizontal / vertical logical links
    arrow(draw, (695, 455), (755, 455), "#3F7FBF")
    arrow(draw, (1360, 455), (1420, 470), "#21897E")
    arrow(draw, (1360, 805), (1420, 830), "#21897E")
    arrow(draw, (1360, 1158), (1420, 1158), "#21897E")
    arrow(draw, (695, 810), (755, 810), "#3F7FBF")
    arrow(draw, (2025, 830), (2085, 830), "#C77B1A")
    arrow(draw, (2025, 1158), (2085, 1158), "#C77B1A")
    arrow(draw, (1055, 590), (1055, 650), "#21897E")
    arrow(draw, (1055, 970), (1055, 1030), "#21897E")

    # The release connector belongs to the background layer: it deliberately
    # disappears beneath the provenance bar rather than cutting through its text.
    arrow(draw, (1395, 1295), (1395, 1575), NAVY, 10)

    # Run/provenance bar
    draw.rounded_rectangle((90, 1360, 2710, 1495), radius=24, fill="#F1F4F7", outline=GRID, width=4)
    draw.text((122, 1384), "Cross-cutting provenance", font=F_NODE, fill=NAVY)
    draw.text(
        (590, 1386),
        "analysis_runs  •  pipeline_failures  •  mapping_remediation_queue  •  arpeggio_attempts / failure classifications",
        font=F_BODY,
        fill=TEXT,
    )

    # Release surfaces
    draw.rounded_rectangle((90, 1575, 2710, 2155), radius=30, fill="#F7FAFC", outline=NAVY, width=5)
    draw.text((130, 1610), "Read-only release surfaces for the V-LiSEMOD v1.0 publication release", font=F_BAND, fill=NAVY)
    draw.text((130, 1665), "Preferred internal V2 query views (version-pinned scientific access)", font=F_NODE, fill="#216E61")
    v2_rows = [
        "v2_structure_context  •  v2_ligand_context  •  v2_ligand_atom_evidence",
        "v2_ligand_comparison_atom_contacts  •  v2_all_chain_lysine_geometry  •  v2_target_lysine_accessibility",
        "v2_protacability_target_context  •  v2_protacability_best  •  v2_attachment_site_candidates",
        "v2_attachment_site_high_priority  •  v2_attachment_site_summary  •  v2_target_browser_ligand_context  •  v2_target_browser_groups",
    ]
    y = 1715
    for line in v2_rows:
        draw.text((155, y), line, font=F_SMALL, fill=TEXT)
        y += 43
    draw.line((130, 1902, 2670, 1902), fill=GRID, width=3)
    draw.text((130, 1932), "Legacy website-compatibility views (retain for prior UI/query behavior; not the authoritative relational contract)", font=F_NODE, fill="#75501B")
    draw.text((155, 1985), "Virus_Proteins  •  Ligand_Atoms_Smiles  •  ligand_atoms  •  RUPLEY_SASA_DATA  •  SMILES_MAP_PDB  •  Functional_Group_Atoms", font=F_SMALL, fill=TEXT)
    draw.text((155, 2028), "Arpeggio_Contacts_Data  •  receptor_binding_pocket  •  solvent_exposed_atoms  •  Covalent_Noncovalent  •  distal_atoms  •  Ligand_Arp_Diagram", font=F_SMALL, fill=TEXT)
    draw.text((155, 2071), "Functional_GROUPED (plus the remaining compatibility projections)", font=F_SMALL, fill=TEXT)
    im.save(PNG, dpi=(300, 300))


def table_name_lines(name, max_chars=26):
    """Wrap underscore-delimited database names without losing the identifier text."""
    parts, lines, line = name.split("_"), [], ""
    for part in parts:
        candidate = f"{line}_{part}" if line else part
        if len(candidate) <= max_chars or not line:
            line = candidate
        else:
            lines.append(line)
            line = part
    if line:
        lines.append(line)
    return lines


def small_node(draw, x0, y0, x1, y1, name, fill, accent):
    draw.rounded_rectangle((x0, y0, x1, y1), radius=18, fill=fill, outline=accent, width=4)
    lines = table_name_lines(name)
    if len(lines) > 2:
        lines = [lines[0], "_".join(lines[1:])]
    y = y0 + (y1 - y0 - len(lines) * 34) // 2 - 2
    for line in lines:
        draw.text((x0 + 16, y), line, font=F_MINI_BOLD, fill=TEXT)
        y += 35


def mini_view(draw, x0, y0, x1, y1, name, accent):
    draw.rounded_rectangle((x0, y0, x1, y1), radius=12, fill="#FFFFFF", outline=accent, width=3)
    text = name if len(name) <= 32 else name.replace("_", "_\n", 1)
    y = y0 + 15 if "\n" not in text else y0 + 5
    draw.multiline_text((x0 + 13, y), text, font=F_SMALL, fill=TEXT, spacing=1)


def build_png_b():
    """Build the expanded, object-level map: every physical table is named."""
    w, h = 2800, 3150
    im = Image.new("RGB", (w, h), BG)
    draw = ImageDraw.Draw(im)
    draw.text((80, 58), "V-LiSEMOD v1.0 expanded physical-table and view map", font=F_TITLE, fill=NAVY)
    draw.text(
        (84, 136),
        "Figure S1.1B. Each colored rectangle is one current physical application table. Arrows show the principal validated logical/provenance paths.",
        font=F_SUB,
        fill=MUTED,
    )

    groups = [
        (70, 225, 695, "STRUCTURE, SYNONYM & PROVENANCE", "#E8F1FB", "#3F7FBF", [
            "structures", "structure_classifications", "structure_context", "analysis_runs", "pipeline_failures", "mapping_remediation_queue", "Ligand_Synonym_Status", "Ligand_Synonyms",
        ]),
        (740, 225, 1365, "LIGAND IDENTITY & ATOMS", "#E9F7F5", "#21897E", [
            "ligands", "ligand_chemistry_sources", "ligand_instances", "ligand_instance_atoms", "ligand_mapping_runs", "ligand_smiles_atom_mapping", "ligand_sasa_atoms", "ligand_atom_geometry",
        ]),
        (1410, 225, 2035, "CONTACT & FUNCTIONAL EVIDENCE", "#FFF4E5", "#C77B1A", [
            "ligand_functional_group_matches", "ligand_functional_group_atoms", "ligand_functional_group_summary", "ligand_arpeggio_runs", "arpeggio_attempts", "arpeggio_derived_atom_map", "arpeggio_failure_classifications", "arpeggio_raw_contact_labels", "arpeggio_unique_atom_pairs", "ligand_binding_pocket_atoms", "receptor_binding_pocket_atoms",
        ]),
        (2080, 225, 2730, "TARGET & PROTACABILITY", "#F4ECFA", "#8751A1", [
            "canonical_ligand_targets", "target_surface_lysines", "target_chain_geometry", "protacability_assessment", "protacability_ligand_inventory", "protacability_lysine_proximity", "protacability_target_context", "protacability_warhead_linkability", "protacability_degrader_readiness", "protacability_attachment_sites", "protacability_attachment_site_summary",
        ]),
    ]
    centers = {}
    for x0, y0, x1, label, fill, accent, names in groups:
        draw.rounded_rectangle((x0, y0, x1, y0 + 55), radius=16, fill=accent)
        draw.text((x0 + 16, y0 + 11), label, font=F_BAND, fill="white")
        y = y0 + 80
        for name in names:
            height = 118
            small_node(draw, x0, y, x1, y + height, name, fill, accent)
            centers[name] = ((x0 + x1) // 2, y + height // 2, x0, x1, y, y + height)
            y += 134

    # Dense but interpretable network of principal joins, drawn behind the release surfaces.
    def edge(a, b, color, source_side="right", target_side="left"):
        sx, sy, x0, x1, y0, y1 = centers[a]
        tx, ty, tx0, tx1, ty0, ty1 = centers[b]
        start = (x1 if source_side == "right" else x0, sy)
        end = (tx0 if target_side == "left" else tx1, ty)
        arrow(draw, start, end, color, 5)

    edge("structures", "ligand_instances", "#3F7FBF")
    edge("structure_context", "ligand_instances", "#3F7FBF")
    edge("ligands", "ligand_instances", "#21897E")
    edge("ligand_chemistry_sources", "ligand_instances", "#21897E")
    edge("ligand_instances", "ligand_instance_atoms", "#21897E", "right", "left")
    edge("ligand_instance_atoms", "ligand_mapping_runs", "#21897E", "right", "left")
    edge("ligand_instance_atoms", "ligand_sasa_atoms", "#21897E", "right", "left")
    edge("ligand_instance_atoms", "ligand_functional_group_atoms", "#21897E", "right", "left")
    edge("ligand_instance_atoms", "ligand_arpeggio_runs", "#21897E", "right", "left")
    edge("ligand_instance_atoms", "ligand_binding_pocket_atoms", "#21897E", "right", "left")
    edge("ligand_instances", "canonical_ligand_targets", "#21897E")
    edge("canonical_ligand_targets", "protacability_target_context", "#8751A1")
    edge("target_surface_lysines", "protacability_lysine_proximity", "#8751A1")
    edge("target_chain_geometry", "protacability_assessment", "#8751A1")
    edge("protacability_target_context", "protacability_attachment_sites", "#8751A1")
    edge("protacability_warhead_linkability", "protacability_degrader_readiness", "#8751A1", "right", "left")
    edge("analysis_runs", "ligand_mapping_runs", "#7C8B98")
    edge("analysis_runs", "ligand_arpeggio_runs", "#7C8B98")

    # Explicit view keys: shown as named release surfaces rather than hiding them in ellipses.
    panel_y = 1815
    draw.rounded_rectangle((70, panel_y, 1365, 3050), radius=28, fill="#F3FBF9", outline="#21897E", width=5)
    draw.rounded_rectangle((1410, panel_y, 2730, 3050), radius=28, fill="#FBF7EE", outline="#A46B17", width=5)
    draw.text((105, panel_y + 32), "13 internal V2 query views", font=F_BAND, fill="#216E61")
    draw.text((1445, panel_y + 32), "13 legacy compatibility views", font=F_BAND, fill="#75501B")
    draw.text((105, panel_y + 80), "Preferred version-pinned scientific access", font=F_SMALL, fill=MUTED)
    draw.text((1445, panel_y + 80), "Prior UI/query support; not the relational contract", font=F_SMALL, fill=MUTED)
    v2 = [
        "v2_structure_context", "v2_ligand_context", "v2_ligand_atom_evidence", "v2_ligand_comparison_atom_contacts",
        "v2_all_chain_lysine_geometry", "v2_target_lysine_accessibility", "v2_protacability_target_context", "v2_protacability_best",
        "v2_attachment_site_candidates", "v2_attachment_site_high_priority", "v2_attachment_site_summary", "v2_target_browser_ligand_context", "v2_target_browser_groups",
    ]
    legacy = [
        "Virus_Proteins", "Ligand_Atoms_Smiles", "Functional_GROUPED", "ligand_atoms", "solvent_exposed_atoms", "RUPLEY_SASA_DATA", "SMILES_MAP_PDB",
        "Functional_Group_Atoms", "Arpeggio_Contacts_Data", "receptor_binding_pocket", "Covalent_Noncovalent", "distal_atoms", "Ligand_Arp_Diagram",
    ]
    for names, x0, x1, accent in ((v2, 105, 1330, "#21897E"), (legacy, 1445, 2695, "#A46B17")):
        col_width = (x1 - x0 - 25) // 2
        for idx, name in enumerate(names):
            col, row = idx % 2, idx // 2
            xx0 = x0 + col * (col_width + 25)
            yy0 = panel_y + 130 + row * 132
            mini_view(draw, xx0, yy0, xx0 + col_width, yy0 + 100, name, accent)

    # Release routes enter the view panels from the evidence and target hubs.
    arrow(draw, (1365, 1320), (1040, panel_y), NAVY, 7)
    arrow(draw, (2035, 1320), (2080, panel_y), NAVY, 7)
    draw.text((78, 3085), "All arrows are logical/provenance routes only; the release validates occurrence traceability rather than declared SQLite foreign keys.", font=F_SMALL, fill=MUTED)
    im.save(PNG_B, dpi=(300, 300))


# The original all-in-one version was deliberately replaced.  A literal 38-table
# spider diagram either becomes illegible or implies relationships that SQLite
# does not declare.  The three panels below show every physical table once,
# anchored to the real key tables and with each card carrying its declared FK
# parents.  This keeps the expanded schema useful at publication scale.
DETAIL_W, DETAIL_H = 2600, 1840
DETAIL_GREY = "#F3F5F7"
DETAIL_GREY_BORDER = "#71808E"


def detail_heading(draw, title, subtitle):
    draw.text((72, 52), title, font=F_TITLE, fill=NAVY)
    draw.text((76, 130), subtitle, font=F_SUB, fill=MUTED)
    draw.rounded_rectangle((72, 188, DETAIL_W - 72, 248), radius=16, fill="#F1F4F7")
    draw.text(
        (95, 203),
        "Solid arrows show selected declared foreign keys (child → referenced parent).  All other declared parents are printed in the cards.",
        font=F_SMALL,
        fill=MUTED,
    )


def detail_band(draw, xy, label, accent):
    x0, y0, x1, y1 = xy
    draw.rounded_rectangle((x0, y0, x1, y1), radius=16, fill=accent)
    draw.text((x0 + 18, y0 + 13), label, font=F_BAND, fill="white")


def detail_card(draw, xy, name, fk, fill, accent, note=None):
    x0, y0, x1, y1 = xy
    draw.rounded_rectangle((x0, y0, x1, y1), radius=18, fill=fill, outline=accent, width=4)
    lines = table_name_lines(name, max_chars=28)
    if len(lines) > 2:
        lines = [lines[0], "_".join(lines[1:])]
    ty = y0 + 15
    for line in lines:
        draw.text((x0 + 18, ty), line, font=F_NODE, fill=TEXT)
        ty += 34
    if note:
        draw.text((x0 + 18, y1 - 53), note, font=F_SMALL, fill=MUTED)
    draw.text((x0 + 18, y1 - 29), fk, font=F_SMALL, fill=accent)


def fk_arrow(draw, start, end, color):
    """Restrained connector used only for the legible core relations."""
    arrow(draw, start, end, color, 6)


def build_png_b1():
    im = Image.new("RGB", (DETAIL_W, DETAIL_H), BG)
    draw = ImageDraw.Draw(im)
    detail_heading(
        draw,
        "Figure S1.1B(i). Expanded schema: identity and common anchors",
        "The core occurrence model. These tables establish the stable identifiers used by the downstream evidence layers.",
    )
    blue_fill, blue = "#E8F1FB", "#3F7FBF"
    teal_fill, teal = "#E9F7F5", "#21897E"
    amber_fill, amber = "#FFF4E5", "#C77B1A"
    grey = "#F3F5F7"

    # Draw the connector layer before *all* visual containers.  This lets cards
    # and section headers mask lines wherever they would overlap text.
    fk_arrow(draw, (420, 575), (420, 540), blue)
    fk_arrow(draw, (1210, 575), (1210, 540), teal)
    fk_arrow(draw, (890, 890), (770, 465), teal)
    fk_arrow(draw, (1210, 1015), (1210, 980), teal)
    fk_arrow(draw, (1830, 605), (1710, 1385), amber)
    fk_arrow(draw, (1830, 835), (1710, 890), amber)

    detail_band(draw, (72, 295, 770, 350), "STRUCTURE & IDENTITY", blue)
    detail_band(draw, (890, 295, 1710, 350), "CANONICAL ANCHOR TABLES", teal)
    detail_band(draw, (1830, 295, 2528, 350), "RUN & QC PROVENANCE", amber)

    # Left source and lookup tables.
    detail_card(draw, (72, 390, 770, 540), "structures", "PK: structure_id", blue_fill, blue, "deposited coordinate entry")
    detail_card(draw, (72, 575, 770, 725), "structure_classifications", "FK → structures", blue_fill, blue, "virus / protein labels")
    detail_card(draw, (72, 760, 770, 910), "structure_context", "no declared FK", blue_fill, blue, "normalized source context")
    detail_card(draw, (72, 945, 770, 1095), "Ligand_Synonyms", "no declared FK", blue_fill, blue, "compatibility identity lookup")
    detail_card(draw, (72, 1130, 770, 1280), "Ligand_Synonym_Status", "no declared FK", blue_fill, blue, "compatibility status lookup")

    # Central actual anchor records.
    detail_card(draw, (890, 390, 1710, 540), "ligands", "PK: ligand_id", teal_fill, teal, "chemical component identity")
    detail_card(draw, (890, 575, 1710, 725), "ligand_chemistry_sources", "FK → ligands", teal_fill, teal, "candidate chemistry provenance")
    detail_card(draw, (890, 800, 1710, 980), "ligand_instances", "FK → structures · ligands", teal_fill, teal, "one deposited ligand occurrence")
    detail_card(draw, (890, 1015, 1710, 1195), "ligand_instance_atoms", "FK → ligand_instances", teal_fill, teal, "occurrence-resolved atom inventory")
    detail_card(draw, (890, 1310, 1710, 1460), "analysis_runs", "PK: run_id", teal_fill, teal, "method / version provenance")

    # Run and QC tables.
    detail_card(draw, (1830, 520, 2528, 690), "pipeline_failures", "FK → run · structure · instance", amber_fill, amber, "explicit processing failures")
    detail_card(draw, (1830, 750, 2528, 920), "mapping_remediation_queue", "FK → instance · first/latest run", amber_fill, amber, "unresolved mapping follow-up")

    draw.rounded_rectangle((72, 1545, 2528, 1740), radius=24, fill="#FAFBFC", outline=DETAIL_GREY_BORDER, width=3)
    draw.text((105, 1575), "How to read the expanded map", font=F_NODE, fill=NAVY)
    draw.text((105, 1625), "The blue/teal anchor tables are the relational backbone.  The remaining panels use their identifiers instead of inventing a direct edge between every derived table.", font=F_BODY, fill=TEXT)
    draw.text((105, 1670), "This is therefore a schema map—not a workflow cartoon—and all 38 current physical tables appear once as colored cards across B(i)–B(iii).", font=F_BODY, fill=TEXT)
    im.save(PNG_B1, dpi=(300, 300))


def build_png_b2():
    im = Image.new("RGB", (DETAIL_W, DETAIL_H), BG)
    draw = ImageDraw.Draw(im)
    detail_heading(
        draw,
        "Figure S1.1B(ii). Expanded schema: atom evidence and interaction analysis",
        "Downstream evidence tables. Each card names its declared parent keys; the three grey cards are repeated anchors from B(i), not additional tables.",
    )
    teal_fill, teal = "#E9F7F5", "#21897E"
    amber_fill, amber = "#FFF4E5", "#C77B1A"
    grey_fill, grey = "#F3F5F7", DETAIL_GREY_BORDER

    detail_band(draw, (72, 295, 2528, 350), "REPEATED KEY ANCHORS FROM PANEL B(i)", DETAIL_GREY_BORDER)
    detail_card(draw, (72, 385, 820, 525), "analysis_runs", "PK: run_id", grey_fill, grey, "repeated anchor")
    detail_card(draw, (925, 385, 1673, 525), "ligand_instances", "PK: ligand_instance_id", grey_fill, grey, "repeated anchor")
    detail_card(draw, (1780, 385, 2528, 525), "ligand_instance_atoms", "PK: ligand_instance_atom_id", grey_fill, grey, "repeated anchor")

    # Draw these first so the cards *and* the colored section bands mask them.
    fk_arrow(draw, (446, 680), (446, 525), teal)
    fk_arrow(draw, (1299, 680), (1299, 525), amber)
    fk_arrow(draw, (2154, 680), (2154, 525), amber)

    detail_band(draw, (72, 590, 820, 645), "MAPPING & ATOM METRICS", teal)
    detail_band(draw, (925, 590, 1673, 645), "ARPEGGIO / CONTACT EVIDENCE", amber)
    detail_band(draw, (1780, 590, 2528, 645), "FUNCTIONAL / POCKET EVIDENCE", amber)

    mapping = [
        ("ligand_mapping_runs", "FK → run · instance"),
        ("ligand_smiles_atom_mapping", "FK → run · instance · atom"),
        ("ligand_sasa_atoms", "FK → run · instance · atom"),
        ("ligand_atom_geometry", "FK → run · instance · atom"),
    ]
    arpeggio = [
        ("ligand_arpeggio_runs", "FK → run · instance"),
        ("arpeggio_attempts", "FK → run · instance"),
        ("arpeggio_derived_atom_map", "FK → run · atom"),
        ("arpeggio_failure_classifications", "FK → run · instance"),
        ("arpeggio_raw_contact_labels", "FK → run · instance · atom"),
        ("arpeggio_unique_atom_pairs", "FK → run · instance · atom"),
        ("receptor_binding_pocket_atoms", "FK → run · instance"),
        ("ligand_binding_pocket_atoms", "FK → run · instance"),
    ]
    functional = [
        ("ligand_functional_group_matches", "FK → run · instance · ligand"),
        ("ligand_functional_group_atoms", "FK → run · instance · atom · match"),
        ("ligand_functional_group_summary", "FK → run · instance"),
    ]
    def stack(items, x0, x1, fill, accent, start=680, step=170):
        for index, (name, fk) in enumerate(items):
            y = start + index * step
            detail_card(draw, (x0, y, x1, y + 145), name, fk, fill, accent)
    stack(mapping, 72, 820, teal_fill, teal, step=190)
    stack(arpeggio, 925, 1673, amber_fill, amber, step=138)
    stack(functional, 1780, 2528, amber_fill, amber, step=210)

    im.save(PNG_B2, dpi=(300, 300))


def build_png_b3():
    im = Image.new("RGB", (DETAIL_W, DETAIL_H), BG)
    draw = ImageDraw.Draw(im)
    detail_heading(
        draw,
        "Figure S1.1B(iii). Expanded schema: target context and PROTACability",
        "Target-associated and attachment-site tables. These eleven tables complete the 38 current physical application tables in the v1.0 release.",
    )
    purple_fill, purple = "#F4ECFA", "#8751A1"
    teal_fill, teal = "#E9F7F5", "#21897E"
    grey_fill, grey = "#F3F5F7", DETAIL_GREY_BORDER

    detail_band(draw, (72, 295, 2528, 350), "REPEATED KEY ANCHORS FROM PANEL B(i)", DETAIL_GREY_BORDER)
    detail_card(draw, (72, 385, 820, 525), "analysis_runs", "PK: run_id", grey_fill, grey, "repeated anchor")
    detail_card(draw, (925, 385, 1673, 525), "ligand_instances", "PK: ligand_instance_id", grey_fill, grey, "repeated anchor")
    detail_card(draw, (1780, 385, 2528, 525), "ligand_instance_atoms", "PK: ligand_instance_atom_id", grey_fill, grey, "repeated anchor")

    # Draw these first so the cards *and* the colored section bands mask them.
    fk_arrow(draw, (446, 695), (446, 525), teal)
    fk_arrow(draw, (1310, 695), (1310, 525), purple)
    fk_arrow(draw, (2141, 695), (2141, 525), purple)

    detail_band(draw, (72, 590, 820, 645), "TARGET CONTEXT & GEOMETRY", teal)
    detail_band(draw, (925, 590, 2528, 645), "PROTACABILITY & ATTACHMENT", purple)
    targets = [
        ("canonical_ligand_targets", "FK → ligand_instance"),
        ("target_chain_geometry", "FK → run · instance"),
        ("target_surface_lysines", "FK → run · instance"),
    ]
    protac_left = [
        ("protacability_ligand_inventory", "FK → run · instance"),
        ("protacability_assessment", "FK → run · instance"),
        ("protacability_target_context", "FK → run · instance"),
        ("protacability_lysine_proximity", "FK → run · instance · atom"),
    ]
    protac_right = [
        ("protacability_warhead_linkability", "FK → run · instance"),
        ("protacability_degrader_readiness", "FK → run · instance"),
        ("protacability_attachment_sites", "FK → run · instance · atom"),
        ("protacability_attachment_site_summary", "FK → run · instance"),
    ]
    for i, (name, fk) in enumerate(targets):
        y = 695 + i * 220
        detail_card(draw, (72, y, 820, y + 170), name, fk, teal_fill, teal)
    for i, (name, fk) in enumerate(protac_left):
        y = 695 + i * 190
        detail_card(draw, (925, y, 1698, y + 155), name, fk, purple_fill, purple)
    for i, (name, fk) in enumerate(protac_right):
        y = 695 + i * 190
        detail_card(draw, (1755, y, 2528, y + 155), name, fk, purple_fill, purple)

    draw.rounded_rectangle((72, 1535, 2528, 1740), radius=24, fill="#FAFBFC", outline=DETAIL_GREY_BORDER, width=3)
    draw.text((105, 1565), "Release views are deliberately not mixed into the physical-table schema", font=F_NODE, fill=NAVY)
    draw.text((105, 1615), "The v1.0 release exposes 13 V2 scientific query views and 13 legacy compatibility views.  They are derived read-only surfaces, not physical tables, and are catalogued separately in the view inventory.", font=F_BODY, fill=TEXT)
    draw.text((105, 1662), "Keeping them separate makes the data model legible and avoids suggesting that a view is a stored data entity.", font=F_BODY, fill=TEXT)
    im.save(PNG_B3, dpi=(300, 300))


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def insert_after(paragraph, new_paragraph):
    paragraph._p.addnext(new_paragraph._p)


def set_font(run, name="Arial", size=10, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii", name)
    run._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi", name)
    run.font.size = Pt(size)
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def compact_legacy_view_table(doc):
    """Remove repeated prose from the compatibility-view table without changing its meaning."""
    table = doc.tables[-1]
    for row in table.rows[1:]:
        paragraph = row.cells[1].paragraphs[0]
        paragraph.text = "Legacy website compatibility; not the v1.0 relational contract."
        for run in paragraph.runs:
            set_font(run, size=9, color=(21, 32, 43))


def insert_map():
    doc = Document(DOCX)
    heading_index = next(
        index for index, p in enumerate(doc.paragraphs)
        if p.text == "S1.4 Logical schema connection map"
    )
    heading = doc.paragraphs[heading_index]
    prior = doc.paragraphs[heading_index + 1]
    next_heading_index = next(
        index for index, p in enumerate(doc.paragraphs[heading_index + 1:], heading_index + 1)
        if p.text == "S1.5 Physical table inventory"
    )
    prior.text = (
        "The diagram maps the current physical application tables by evidence layer and shows the released "
        "read-only surfaces used for reproducible retrieval."
    )
    for run in prior.runs:
        set_font(run, size=10, color=(21, 32, 43))
    heading.paragraph_format.keep_with_next = False
    prior.paragraph_format.keep_with_next = False
    # Replace the old monospace map, or a prior generated map, as one local block.
    for paragraph in doc.paragraphs[heading_index + 2:next_heading_index]:
        remove_paragraph(paragraph)

    figure_a = doc.add_paragraph()
    figure_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_a.paragraph_format.space_before = Pt(8)
    figure_a.paragraph_format.space_after = Pt(3)
    figure_a.paragraph_format.keep_with_next = True
    figure_a.add_run().add_picture(str(PNG), width=Inches(6.55))
    insert_after(prior, figure_a)

    caption_a = doc.add_paragraph()
    caption_a.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption_a.paragraph_format.space_before = Pt(2)
    caption_a.paragraph_format.space_after = Pt(8)
    run = caption_a.add_run(
        "Figure S1.1A. Layer-level V-LiSEMOD v1.0 schema connection map. Rounded rectangles identify current "
        "physical tables or table families; the lower panel names the version-pinned V2 views and legacy "
        "compatibility views."
    )
    set_font(run, size=9, color=(70, 80, 90), italic=True)
    insert_after(figure_a, caption_a)

    current = caption_a
    panels = [
        (PNG_B1,
         "Figure S1.1B(i). Expanded V-LiSEMOD v1.0 schema: identity and common anchors. Solid arrows identify "
         "the small set of core declared foreign-key relationships; every card lists its own declared parent keys."),
        (PNG_B2,
         "Figure S1.1B(ii). Expanded V-LiSEMOD v1.0 schema: atom evidence and interaction analysis. The repeated "
         "grey cards are key anchors from panel (i), not additional physical tables."),
        (PNG_B3,
         "Figure S1.1B(iii). Expanded V-LiSEMOD v1.0 schema: target context and PROTACability. Together, panels "
         "(i)–(iii) show all 38 current physical application tables once as colored cards; grey cards are repeated key anchors."),
    ]
    for image_path, caption_text in panels:
        figure = doc.add_paragraph()
        figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure.paragraph_format.page_break_before = True
        figure.paragraph_format.space_before = Pt(0)
        figure.paragraph_format.space_after = Pt(3)
        figure.paragraph_format.keep_with_next = True
        figure.add_run().add_picture(str(image_path), width=Inches(6.55))
        insert_after(current, figure)

        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(8)
        run = caption.add_run(caption_text)
        set_font(run, size=9, color=(70, 80, 90), italic=True)
        insert_after(figure, caption)
        current = caption
    compact_legacy_view_table(doc)
    doc.save(DOCX)


if __name__ == "__main__":
    build_png()
    build_png_b1()
    build_png_b2()
    build_png_b3()
    insert_map()
